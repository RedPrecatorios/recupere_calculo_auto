from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import structlog
from openpyxl import load_workbook

from recupere_monday.config import Settings
from recupere_monday.extraction.demonstrativo_fields import extract_campos_from_pdf
from recupere_monday.llm import build_pdf_extractor

logger = structlog.get_logger(__name__)


@dataclass
class CalculoIATestResult:
    processed: int = 0
    skipped: int = 0
    failed: int = 0
    requerente_key: str | None = None
    output_path: Path | None = None
    errors: list[str] = field(default_factory=list)


class CalculoIATestOrchestrator:
    """Faz uma chamada real de teste enviando apenas PDFs à IA."""

    def __init__(self, settings: Settings, *, validar_leitura: bool = False) -> None:
        self._settings = settings
        self._input_dir = settings.requerentes_output_path
        self._validar_leitura = validar_leitura

    def run(self) -> CalculoIATestResult:
        result = CalculoIATestResult()
        for payload_path in sorted(self._input_dir.rglob("payload_calculo.json")):
            payload = _read_json(payload_path)
            attachment_paths = _eligible_pdf_paths(payload)
            if not attachment_paths:
                result.skipped += 1
                continue
            fluxo = _resolve_calculo_fluxo(payload, self._settings)
            llm = _build_llm_processor(
                self._settings,
                validar_leitura=self._validar_leitura,
                fluxo=fluxo,
            )

            try:
                extraction = llm.extract_from_pdfs(attachment_paths)
                output_path = payload_path.parent / (
                    "validacao_leitura_pdf_ia.json"
                    if self._validar_leitura
                    else "resultado_calculo_teste.json"
                )
                output_path.write_text(
                    json.dumps(
                        {
                            "tipo_teste": (
                                "validacao_leitura_pdf"
                                if self._validar_leitura
                                else "calculo"
                            ),
                            "requerente": payload.get("requerente"),
                            "arquivos_enviados": [
                                str(path) for path in attachment_paths
                            ],
                            "pdfs_enviados": [
                                str(path)
                                for path in attachment_paths
                                if path.suffix.lower() == ".pdf"
                            ],
                            "planilha_modelo_enviada": (
                                None
                            ),
                            "planilha_modelo_referenciada_no_prompt": (
                                str(self._settings.planilha_juros_modelo_path)
                                if (
                                    not self._validar_leitura
                                    and fluxo.get("tipo_fluxo")
                                    == "juros_zerado_usar_planilha"
                                )
                                else None
                            ),
                            "decisao_fluxo": fluxo,
                            "resultado": extraction.parsed_data,
                        },
                        ensure_ascii=False,
                        indent=2,
                    ),
                    encoding="utf-8",
                )
                result.processed = 1
                result.requerente_key = str(payload.get("requerente_key"))
                result.output_path = output_path
                logger.info(
                    "calculo_ia_test_success",
                    requerente=result.requerente_key,
                    arquivos=[path.name for path in attachment_paths],
                    fluxo=fluxo.get("tipo_fluxo"),
                    output=str(output_path),
                )
                return result
            except Exception as exc:  # noqa: BLE001
                result.failed += 1
                result.errors.append(f"{payload_path}: {exc}")
                logger.exception(
                    "calculo_ia_test_failed",
                    payload=str(payload_path),
                    error=str(exc),
                )
                return result

        result.errors.append("Nenhum requerente elegível com IR retido encontrado")
        return result


def _eligible_pdf_paths(payload: dict[str, Any]) -> list[Path]:
    pdfs: list[Path] = []
    for registro in payload.get("registros", []):
        status = registro.get("status", {})
        if not status.get("pronto_para_calculo") or not status.get("tem_ir_retido"):
            continue

        demonstrativo_pdf = (
            registro.get("demonstrativo", {}).get("pdf_extraido_path")
        )
        comprovante_pdf = registro.get("comprovante", {}).get("pdf_path")
        if demonstrativo_pdf and comprovante_pdf:
            pdfs.extend([Path(demonstrativo_pdf), Path(comprovante_pdf)])
    return _dedupe_paths(pdfs)


def _dedupe_paths(paths: list[Path]) -> list[Path]:
    unique: list[Path] = []
    seen: set[str] = set()
    for path in paths:
        key = str(path)
        if key in seen:
            continue
        seen.add(key)
        unique.append(path)
    return unique


def _with_spreadsheet_context(prompt: str, spreadsheet_path: Path) -> str:
    if not spreadsheet_path.exists():
        logger.warning(
            "planilha_juros_modelo_not_found",
            path=str(spreadsheet_path),
        )
        return prompt
    return (
        f"{prompt}\n\n"
        "Referência textual extraída da planilha Excel modelo do escritório:\n"
        f"{_spreadsheet_reference_text(spreadsheet_path)}"
    )


def _spreadsheet_reference_text(path: Path, *, max_cells: int = 500) -> str:
    workbook = load_workbook(path, data_only=False, read_only=True)
    lines = [f"Arquivo: {path.name}"]
    cells_count = 0
    for sheet in workbook.worksheets:
        lines.append(f"Aba: {sheet.title}")
        for row in sheet.iter_rows():
            values = []
            for cell in row:
                if cell.value is None:
                    continue
                values.append(f"{cell.coordinate}={cell.value}")
            if not values:
                continue
            lines.append(" | ".join(values))
            cells_count += len(values)
            if cells_count >= max_cells:
                lines.append("Conteúdo truncado para manter o prompt compacto.")
                return "\n".join(lines)
    return "\n".join(lines)


def _build_llm_processor(
    settings: Settings,
    *,
    validar_leitura: bool,
    fluxo: dict[str, Any],
):
    if validar_leitura:
        prompt = settings.resolved_validacao_leitura_pdf_prompt()
    elif fluxo.get("tipo_fluxo") == "juros_zerado_usar_planilha":
        prompt = _with_spreadsheet_context(
            settings.resolved_calculo_juros_zerado_prompt(),
            settings.planilha_juros_modelo_path,
        )
    else:
        prompt = settings.resolved_calculo_rra_direto_prompt()
    return build_pdf_extractor(settings, prompt)


def _resolve_calculo_fluxo(
    payload: dict[str, Any],
    settings: Settings,
) -> dict[str, Any]:
    demonstrativo_pdf = _first_demonstrativo_pdf(payload)
    juros_mora_selic: float | None = None
    if demonstrativo_pdf and demonstrativo_pdf.exists():
        juros_mora_selic = extract_campos_from_pdf(demonstrativo_pdf).juros_mora_selic

    if juros_mora_selic is not None and abs(juros_mora_selic) < 0.005:
        tipo_fluxo = "juros_zerado_usar_planilha"
        prompt_path = settings.openai_calculo_juros_zerado_prompt_file
    else:
        tipo_fluxo = "rra_direto_juros_mora_selic_nao_zerado"
        prompt_path = settings.openai_calculo_rra_direto_prompt_file

    return {
        "tipo_fluxo": tipo_fluxo,
        "juros_mora_selic_detectado": juros_mora_selic,
        "demonstrativo_pdf_usado_na_decisao": (
            str(demonstrativo_pdf) if demonstrativo_pdf else None
        ),
        "prompt_usado": str(prompt_path),
    }


def _first_demonstrativo_pdf(payload: dict[str, Any]) -> Path | None:
    for registro in payload.get("registros", []):
        demonstrativo_pdf = registro.get("demonstrativo", {}).get("pdf_extraido_path")
        if demonstrativo_pdf:
            return Path(demonstrativo_pdf)
    return None


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))
