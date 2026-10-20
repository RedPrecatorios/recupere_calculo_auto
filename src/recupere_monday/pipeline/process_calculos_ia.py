from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import structlog
import fitz
from openpyxl import load_workbook

from recupere_monday.config import Settings
from recupere_monday.extraction.demonstrativo_fields import extract_campos_from_pdf
from recupere_monday.extraction.numero_meses_rra import escolher_numero_meses_rra
from recupere_monday.llm import build_pdf_extractor
from recupere_monday.monday.client import MondayClient
from recupere_monday.receita.rra import ReceitaRRAClient, RRAInput
from recupere_monday.receita.selic import (
    SicalcSelicClient,
    calcular_imposto_indevido_atualizado,
)
from recupere_monday.openai.processor import parse_llm_json
from recupere_monday.storage.operational_db import OperationalDBRepository

logger = structlog.get_logger(__name__)

OPENAI_TOTAL_FILE_LIMIT_BYTES = 50 * 1024 * 1024


@dataclass
class CalculosIAResult:
    total_registros: int = 0
    elegiveis: int = 0
    processed: int = 0
    skipped: int = 0
    failed: int = 0
    monday_encaminhados: int = 0
    monday_ja_encaminhados: int = 0
    monday_nao_elegiveis: int = 0
    monday_failed: int = 0
    output_summary_path: Path | None = None
    errors: list[str] = field(default_factory=list)


class CalculosIAOrchestrator:
    """Processa todos os registros elegíveis, um caso por chamada à IA."""

    def __init__(self, settings: Settings) -> None:
        self._settings = settings
        self._input_dir = settings.requerentes_output_path
        self._operational_db = OperationalDBRepository.from_settings(settings)

    def run(
        self,
        *,
        limit: int | None = None,
        process_filter: str | None = None,
        force_reprocess: bool = False,
    ) -> CalculosIAResult:
        self._force_reprocess = force_reprocess
        result = CalculosIAResult()
        summary_rows: list[dict[str, Any]] = []

        for payload_path in sorted(self._input_dir.rglob("payload_calculo.json")):
            payload = _read_json(payload_path)
            for registro in payload.get("registros", []):
                if (
                    process_filter
                    and registro.get("demonstrativo", {}).get("processo_pasta")
                    != process_filter
                ):
                    continue
                result.total_registros += 1
                record_id = _record_output_id(registro)
                blockers = _eligibility_blockers(registro)
                if blockers:
                    result.skipped += 1
                    summary_rows.append(
                        {
                            "status": "pulado",
                            "requerente_key": payload.get("requerente_key"),
                            "record_id": record_id,
                            "motivo": "registro_nao_elegivel",
                            "detalhes": blockers,
                            "output_path": None,
                            "erro": None,
                        }
                    )
                    continue
                result.elegiveis += 1
                if (
                    self._operational_db
                    and self._operational_db.should_skip_calculo(registro)
                    and not getattr(self, "_force_reprocess", False)
                ):
                    result.skipped += 1
                    summary_rows.append(
                        {
                            "status": "pulado",
                            "requerente_key": payload.get("requerente_key"),
                            "record_id": record_id,
                            "motivo": "ja_processado_no_mysql",
                            "detalhes": [
                                "ar_calculos_ia.status=processado para esta requisicao"
                            ],
                            "output_path": None,
                            "erro": None,
                        }
                    )
                    logger.info(
                        "calculo_ia_skipped_already_processed",
                        record_id=record_id,
                    )
                    continue
                if limit is not None and result.processed >= limit:
                    result.skipped += 1
                    summary_rows.append(
                        {
                            "status": "pulado",
                            "requerente_key": payload.get("requerente_key"),
                            "record_id": record_id,
                            "motivo": "limite_atingido",
                            "detalhes": [f"limit={limit}"],
                            "output_path": None,
                            "erro": None,
                        }
                    )
                    continue

                output_path = payload_path.parent / "calculos_ia" / f"{record_id}.json"
                try:
                    row = self._process_record(payload, registro, output_path)
                    result.processed += 1
                    monday_status = (
                        row.get("monday_encaminhamento") or {}
                    ).get("status")
                    if monday_status == "encaminhado":
                        result.monday_encaminhados += 1
                    elif monday_status == "ja_encaminhado":
                        result.monday_ja_encaminhados += 1
                    elif monday_status == "nao_elegivel":
                        result.monday_nao_elegiveis += 1
                    elif monday_status == "falha":
                        result.monday_failed += 1
                    summary_rows.append(row)
                except Exception as exc:  # noqa: BLE001
                    result.failed += 1
                    detail = f"{payload_path} | {record_id}: {exc}"
                    result.errors.append(detail)
                    summary_rows.append(
                        {
                            "status": "falha",
                            "requerente_key": payload.get("requerente_key"),
                            "record_id": record_id,
                            "output_path": str(output_path),
                            "erro": str(exc),
                            "error_class": "unknown",
                        }
                    )
                    if self._operational_db:
                        self._operational_db.persist_calculation_failure(
                            registro=registro,
                            error_message=str(exc),
                            error_class="unknown",
                            output_path=output_path,
                        )
                    logger.exception(
                        "calculo_ia_record_failed",
                        payload=str(payload_path),
                        record_id=record_id,
                        error=str(exc),
                    )

        result.output_summary_path = self._write_summary(summary_rows)
        return result

    def _process_record(
        self,
        payload: dict[str, Any],
        registro: dict[str, Any],
        output_path: Path,
    ) -> dict[str, Any]:
        demonstrativo_pdf = Path(registro["demonstrativo"]["pdf_extraido_path"])
        comprovante_pdf = Path(registro["comprovante"]["pdf_path"])
        original_attachments = [demonstrativo_pdf, comprovante_pdf]
        attachments = original_attachments
        total_size = _total_size(attachments)
        fluxo = _resolve_fluxo(
            demonstrativo_pdf,
            self._settings,
        )

        if total_size > OPENAI_TOTAL_FILE_LIMIT_BYTES:
            attachments = _optimize_attachments_for_openai(
                original_attachments,
                output_path.parent / "_arquivos_otimizados" / output_path.stem,
            )
            total_size = _total_size(attachments)
            if total_size > OPENAI_TOTAL_FILE_LIMIT_BYTES:
                raise ValueError(
                    "Arquivos excedem 50MB na mesma chamada mesmo após otimização: "
                    f"{round(total_size / (1024 * 1024), 2)}MB"
                )
            logger.info(
                "calculo_ia_attachments_optimized",
                record_id=registro.get("id"),
                original_mb=round(_total_size(original_attachments) / (1024 * 1024), 2),
                optimized_mb=round(total_size / (1024 * 1024), 2),
            )

        openai = _build_llm_processor(self._settings, fluxo)
        extraction = openai.extract_from_pdfs(attachments)
        parsed_data = _normalize_parsed_extraction(extraction.parsed_data)
        ajustes_deterministicos = _apply_comprovante_deterministic_overrides(
            parsed_data,
            registro,
            demonstrativo_pdf,
        )
        resultado_receita_rra = _calculate_receita_rra(
            parsed_data,
            timeout_seconds=self._settings.request_timeout_seconds,
        )
        resultado_imposto_atualizado = _calculate_imposto_indevido_atualizado(
            parsed_data,
            resultado_receita_rra,
            timeout_seconds=self._settings.request_timeout_seconds,
        )
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_payload = {
            "tipo_processamento": "calculo_ia_registro",
            "requerente": payload.get("requerente"),
            "requerente_key": payload.get("requerente_key"),
            "registro_id": registro.get("id"),
            "arquivos_originais": [str(path) for path in original_attachments],
            "arquivos_enviados": [str(path) for path in attachments],
            "arquivos_otimizados_para_openai": attachments != original_attachments,
            "tamanho_total_arquivos_bytes": total_size,
            "decisao_fluxo": fluxo,
            "planilha_modelo_referenciada_no_prompt": (
                str(self._settings.planilha_juros_modelo_path)
                if fluxo.get("tipo_fluxo")
                in {
                    "juros_zerado_usar_planilha",
                    "juros_indeterminado_usar_planilha_por_segurança",
                }
                else None
            ),
            "resultado": parsed_data,
            "ajustes_deterministicos": ajustes_deterministicos,
            "resultado_receita_rra": resultado_receita_rra,
            "resultado_imposto_retido_indevidamente_atualizado": (
                resultado_imposto_atualizado
            ),
        }
        output_payload["monday_encaminhamento"] = self._encaminhar_para_monday(
            output_payload,
            registro,
            output_path,
        )
        output_path.write_text(
            json.dumps(output_payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        if self._operational_db:
            self._operational_db.persist_calculation_result(
                output_payload=output_payload,
                output_path=output_path,
                registro=registro,
                openai_model=self._settings.active_llm_model,
            )
        logger.info(
            "calculo_ia_record_success",
            requerente=payload.get("requerente_key"),
            record_id=registro.get("id"),
            fluxo=fluxo.get("tipo_fluxo"),
            receita_rra_status=resultado_receita_rra.get("status"),
            imposto_atualizado_status=resultado_imposto_atualizado.get("status"),
            output=str(output_path),
        )
        return {
            "status": "processado",
            "requerente_key": payload.get("requerente_key"),
            "record_id": registro.get("id"),
            "fluxo": fluxo.get("tipo_fluxo"),
            "receita_rra_status": resultado_receita_rra.get("status"),
            "receita_rra_motivo": resultado_receita_rra.get("motivo"),
            "receita_rra_campos_faltantes": resultado_receita_rra.get(
                "campos_faltantes"
            ),
            "imposto_atualizado_status": resultado_imposto_atualizado.get("status"),
            "imposto_atualizado_motivo": resultado_imposto_atualizado.get("motivo"),
            "monday_encaminhamento": output_payload.get("monday_encaminhamento"),
            "output_path": str(output_path),
            "erro": None,
        }

    def _encaminhar_para_monday(
        self,
        output_payload: dict[str, Any],
        registro: dict[str, Any],
        output_path: Path,
    ) -> dict[str, Any]:
        elegibilidade = _monday_forwarding_eligibility(output_payload)
        if not elegibilidade["elegivel"]:
            return {
                "status": "nao_elegivel",
                "motivo": elegibilidade["motivo"],
            }

        registro_id = str(output_payload.get("registro_id") or registro.get("id") or "")
        board_id = self._settings.monday_board_id

        # 1) Fonte de verdade: MySQL ar_monday_envios
        if self._operational_db and registro_id:
            existing_db = self._operational_db.find_monday_envio(
                registro_id=registro_id,
                monday_board_id=board_id,
            )
            if (
                existing_db
                and existing_db.get("status") == "enviado"
                and existing_db.get("monday_item_id")
            ):
                return {
                    "status": "ja_encaminhado",
                    "motivo": "item_monday_ja_registrado_no_mysql",
                    "item_id": existing_db.get("monday_item_id"),
                    "item_name": existing_db.get("item_name"),
                    "monday_envio_id": existing_db.get("monday_envio_id"),
                }

        # 2) Fallbacks locais (legado / DB desligado)
        previous = _read_existing_monday_forwarding(output_path)
        if previous.get("item_id"):
            return {
                "status": "ja_encaminhado",
                "motivo": "item_monday_ja_registrado_no_json_anterior",
                "item_id": previous.get("item_id"),
                "item_name": previous.get("item_name"),
                "monday_envio_id": previous.get("monday_envio_id"),
            }

        state = _read_monday_envio_state(self._settings.monday_envio_state_path)
        state_record = (state.get("registros") or {}).get(registro_id)
        if isinstance(state_record, dict) and state_record.get("item_id"):
            return {
                "status": "ja_encaminhado",
                "motivo": "item_monday_ja_registrado_no_controle_local",
                "item_id": state_record.get("item_id"),
                "item_name": state_record.get("item_name"),
                "monday_envio_id": state_record.get("monday_envio_id"),
            }

        if self._operational_db:
            monday_envio_id = self._operational_db.next_monday_envio_id(
                monday_board_id=board_id,
                configured_last_id=self._settings.monday_ultimo_envio_id,
            )
        else:
            monday_envio_id = _next_monday_envio_id(
                state,
                self._settings.monday_ultimo_envio_id,
            )

        monday_payload = _build_monday_payload(
            output_payload,
            registro,
            output_path,
            monday_envio_id=monday_envio_id,
        )

        # 3) Última barreira: o próprio board Monday
        duplicado = self._find_monday_duplicate(monday_payload["campos"])
        if duplicado:
            item_id_existente = str(duplicado.get("id") or "")
            item_name_existente = str(duplicado.get("name") or "")
            if self._operational_db and registro_id and item_id_existente:
                try:
                    self._operational_db.record_monday_duplicado(
                        registro_id=registro_id,
                        monday_board_id=board_id,
                        monday_group_id=self._settings.monday_group_id,
                        monday_item_id=item_id_existente,
                        item_name=item_name_existente,
                        registro=registro,
                        campos=monday_payload.get("campos"),
                    )
                except Exception as exc:  # noqa: BLE001
                    logger.warning(
                        "calculo_ia_monday_duplicado_registro_falhou",
                        registro_id=registro_id,
                        error=str(exc),
                    )
            logger.info(
                "calculo_ia_monday_duplicado",
                registro_id=registro_id,
                item_id=item_id_existente,
                item_name=item_name_existente,
            )
            return {
                "status": "ja_encaminhado",
                "motivo": "item_ja_existe_no_board_monday",
                "item_id": item_id_existente,
                "item_name": item_name_existente,
            }

        if self._operational_db and registro_id:
            try:
                reserved = self._operational_db.reserve_monday_envio(
                    registro_id=registro_id,
                    monday_board_id=board_id,
                    monday_group_id=self._settings.monday_group_id,
                    monday_envio_id=monday_envio_id,
                    registro=registro,
                    item_name=monday_payload["item_name"],
                    campos=monday_payload.get("campos"),
                )
                if (
                    reserved.get("status") == "enviado"
                    and reserved.get("monday_item_id")
                ):
                    return {
                        "status": "ja_encaminhado",
                        "motivo": "item_monday_ja_registrado_no_mysql",
                        "item_id": reserved.get("monday_item_id"),
                        "item_name": reserved.get("item_name"),
                        "monday_envio_id": reserved.get("monday_envio_id"),
                    }
                if reserved.get("monday_envio_id"):
                    monday_envio_id = int(reserved["monday_envio_id"])
            except Exception as exc:  # noqa: BLE001
                logger.warning(
                    "calculo_ia_monday_reserve_failed",
                    registro_id=registro_id,
                    error=str(exc),
                )

        attached_files: list[dict[str, str]] = []
        try:
            with MondayClient(
                api_token=self._settings.monday_api_token,
                board_id=self._settings.monday_board_id,
                group_id=self._settings.monday_group_id,
                timeout_seconds=self._settings.request_timeout_seconds,
            ) as monday:
                item_id = monday.create_item(
                    monday_payload["item_name"],
                    MondayClient.build_column_values(
                        monday_payload["campos"],
                        self._settings.parse_monday_mapping(),
                    ),
                )
                if self._settings.monday_attach_pdfs:
                    attached_files = _attach_monday_pdfs(
                        monday,
                        item_id=item_id,
                        settings=self._settings,
                        monday_payload=monday_payload,
                    )
        except Exception as exc:  # noqa: BLE001
            logger.exception(
                "calculo_ia_monday_forward_failed",
                record_id=registro.get("id"),
                output=str(output_path),
                error=str(exc),
            )
            if self._operational_db and registro_id:
                try:
                    self._operational_db.mark_monday_falha(
                        registro_id=registro_id,
                        monday_board_id=board_id,
                        error_message=str(exc),
                    )
                except Exception as mark_exc:  # noqa: BLE001
                    logger.warning(
                        "calculo_ia_monday_mark_falha_failed",
                        error=str(mark_exc),
                    )
            return {
                "status": "falha",
                "motivo": str(exc),
                "item_name": monday_payload["item_name"],
                "monday_envio_id": monday_envio_id,
            }

        if self._operational_db and registro_id:
            try:
                self._operational_db.mark_monday_enviado(
                    registro_id=registro_id,
                    monday_board_id=board_id,
                    monday_item_id=str(item_id),
                    monday_envio_id=monday_envio_id,
                    item_name=monday_payload["item_name"],
                )
            except Exception as exc:  # noqa: BLE001
                logger.warning(
                    "calculo_ia_monday_mark_enviado_failed",
                    registro_id=registro_id,
                    item_id=item_id,
                    error=str(exc),
                )

        _write_monday_envio_state(
            self._settings.monday_envio_state_path,
            state,
            registro_id=registro_id,
            monday_envio_id=monday_envio_id,
            item_id=item_id,
            item_name=monday_payload["item_name"],
            output_path=output_path,
        )
        logger.info(
            "calculo_ia_monday_forwarded",
            record_id=registro.get("id"),
            item_id=item_id,
            monday_envio_id=monday_envio_id,
            arquivos_anexados=len(attached_files),
            output=str(output_path),
        )
        return {
            "status": "encaminhado",
            "motivo": "calculo_efetuado_com_imposto_indevido_a_restituir",
            "item_id": item_id,
            "item_name": monday_payload["item_name"],
            "monday_envio_id": monday_envio_id,
            "campos_disponiveis": sorted(monday_payload["campos"].keys()),
            "arquivos_anexados": attached_files,
        }

    def _find_monday_duplicate(
        self,
        campos: dict[str, Any],
    ) -> dict[str, Any] | None:
        if not self._settings.monday_dedupe_enabled:
            return None
        fields = self._settings.parse_monday_dedupe_fields()
        columns = _build_monday_dedupe_columns(
            campos,
            self._settings.parse_monday_mapping(),
            fields,
        )
        # Chave incompleta encontraria itens de outras requisições do mesmo
        # requerente, então a checagem remota é ignorada.
        if not columns or len(columns) != len(fields):
            logger.warning(
                "calculo_ia_monday_dedupe_sem_chave",
                campos=self._settings.monday_dedupe_fields,
                resolvidos=sorted(columns),
            )
            return None
        try:
            with MondayClient(
                api_token=self._settings.monday_api_token,
                board_id=self._settings.monday_board_id,
                group_id=self._settings.monday_group_id,
                timeout_seconds=self._settings.request_timeout_seconds,
            ) as monday:
                return monday.find_existing_item(columns)
        except Exception as exc:  # noqa: BLE001
            # Falha na consulta não pode bloquear o encaminhamento.
            logger.warning(
                "calculo_ia_monday_dedupe_falhou",
                error=str(exc),
            )
            return None

    def _write_summary(self, rows: list[dict[str, Any]]) -> Path:
        output_path = self._input_dir / "_calculos_ia_summary.json"
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(
            json.dumps(
                {
                    "summary": {
                        "total_linhas": len(rows),
                        "processados": sum(
                            1 for row in rows if row.get("status") == "processado"
                        ),
                        "falhas": sum(
                            1 for row in rows if row.get("status") == "falha"
                        ),
                        "pulados": sum(
                            1 for row in rows if row.get("status") == "pulado"
                        ),
                        "rra_calculados": sum(
                            1
                            for row in rows
                            if row.get("receita_rra_status") == "calculado"
                        ),
                        "rra_nao_calculados": sum(
                            1
                            for row in rows
                            if row.get("receita_rra_status") == "nao_calculado"
                        ),
                        "imposto_atualizado_calculados": sum(
                            1
                            for row in rows
                            if row.get("imposto_atualizado_status") == "calculado"
                        ),
                        "imposto_atualizado_nao_calculados": sum(
                            1
                            for row in rows
                            if row.get("imposto_atualizado_status") == "nao_calculado"
                        ),
                        "monday_encaminhados": sum(
                            1
                            for row in rows
                            if (row.get("monday_encaminhamento") or {}).get("status")
                            == "encaminhado"
                        ),
                        "monday_ja_encaminhados": sum(
                            1
                            for row in rows
                            if (row.get("monday_encaminhamento") or {}).get("status")
                            == "ja_encaminhado"
                        ),
                        "monday_nao_elegiveis": sum(
                            1
                            for row in rows
                            if (row.get("monday_encaminhamento") or {}).get("status")
                            == "nao_elegivel"
                        ),
                        "monday_falhas": sum(
                            1
                            for row in rows
                            if (row.get("monday_encaminhamento") or {}).get("status")
                            == "falha"
                        ),
                    },
                    "pendencias_calculo": [
                        row
                        for row in rows
                        if row.get("receita_rra_status") == "nao_calculado"
                        or row.get("imposto_atualizado_status") == "nao_calculado"
                    ],
                    "registros": rows,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        return output_path


def _is_eligible(registro: dict[str, Any]) -> bool:
    return not _eligibility_blockers(registro)


def _monday_forwarding_eligibility(output_payload: dict[str, Any]) -> dict[str, Any]:
    receita_rra = output_payload.get("resultado_receita_rra") or {}
    imposto_atualizado = (
        output_payload.get("resultado_imposto_retido_indevidamente_atualizado") or {}
    )
    if receita_rra.get("status") != "calculado":
        return {
            "elegivel": False,
            "motivo": "receita_rra_nao_calculada",
        }
    if imposto_atualizado.get("status") != "calculado":
        return {
            "elegivel": False,
            "motivo": "imposto_atualizado_nao_calculado",
        }
    if _decimal(imposto_atualizado.get("imposto_indevido")) <= Decimal("0"):
        return {
            "elegivel": False,
            "motivo": "sem_imposto_indevido_a_restituir",
        }
    return {
        "elegivel": True,
        "motivo": "calculo_efetuado_com_imposto_indevido_a_restituir",
    }


def _read_existing_monday_forwarding(output_path: Path) -> dict[str, Any]:
    if not output_path.exists():
        return {}
    try:
        payload = _read_json(output_path)
    except Exception:  # noqa: BLE001
        return {}
    forwarding = payload.get("monday_encaminhamento")
    return forwarding if isinstance(forwarding, dict) else {}


def _attach_monday_pdfs(
    monday: MondayClient,
    *,
    item_id: int,
    settings: Settings,
    monday_payload: dict[str, Any],
) -> list[dict[str, str]]:
    campos = monday_payload.get("campos") or {}
    pdf_paths: list[Path] = []
    for key in ("demonstrativo_pdf", "comprovante_pdf"):
        raw = campos.get(key)
        if raw:
            pdf_paths.append(Path(str(raw)))
    if not pdf_paths:
        logger.warning(
            "monday_attach_pdfs_skipped",
            item_id=item_id,
            motivo="pdfs_ausentes_no_payload",
        )
        return []

    try:
        return monday.attach_pdfs_to_files_column(
            item_id,
            pdf_paths,
            column_id=settings.monday_files_column_id,
        )
    except Exception as exc:  # noqa: BLE001
        # Item já criado: não falha o encaminhamento por erro de anexo.
        logger.exception(
            "monday_attach_pdfs_failed",
            item_id=item_id,
            error=str(exc),
        )
        return []


def _build_monday_dedupe_columns(
    campos: dict[str, Any],
    mapping: dict[str, str],
    fields: tuple[str, ...],
) -> dict[str, str]:
    columns: dict[str, str] = {}
    for field in fields:
        column_id = mapping.get(field)
        value = campos.get(field)
        if not column_id or value in (None, ""):
            continue
        columns[column_id] = str(value)
    return columns


def _build_monday_payload(
    output_payload: dict[str, Any],
    registro: dict[str, Any],
    output_path: Path,
    *,
    monday_envio_id: int,
) -> dict[str, Any]:
    case = _get_first_case(output_payload.get("resultado") or {})
    identificacao = case.get("identificacao") or {}
    dados_comprovante = case.get("dados_comprovante_bancario") or {}
    receita_rra = output_payload.get("resultado_receita_rra") or {}
    receita_resultado = receita_rra.get("resultado") or {}
    imposto_atualizado = (
        output_payload.get("resultado_imposto_retido_indevidamente_atualizado") or {}
    )
    selic = imposto_atualizado.get("selic") or {}
    demonstrativo = registro.get("demonstrativo") or {}
    comprovante = registro.get("comprovante") or {}
    demonstrativo_pdf = _first(output_payload, "arquivos_originais")
    demonstrativo_metadata = _resolve_demonstrativo_metadata(
        demonstrativo,
        demonstrativo_pdf,
    )
    data_levantamento = (
        dados_comprovante.get("data_levantamento")
        or comprovante.get("data_levantamento")
    )
    demonstrativo_financials = (
        _extract_demonstrativo_receita_financials(Path(str(demonstrativo_pdf)))
        if demonstrativo_pdf
        else {}
    )
    valor_total_pagamento = _money_to_float(
        (case.get("dados_demonstrativo") or {}).get("valor_total_pagamento")
        or demonstrativo_financials.get("valor_total_devido_pagamento")
        or (receita_rra.get("entrada") or {}).get("rendimentos_recebidos")
    )
    imposto_retido_indevidamente_atualizado = _money_to_float(
        imposto_atualizado.get("imposto_retido_indevidamente_atualizado")
    )
    # IR Retido na Monday = IR localizado no comprovante de levantamento.
    ir_retido_comprovante = _money_to_float(
        (comprovante.get("ir_retido") or {}).get("valor_ir_retido")
        or dados_comprovante.get("imposto_retido_na_fonte")
        or imposto_atualizado.get("imposto_retido")
    )

    campos = {
        "monday_envio_id": monday_envio_id,
        "elemento": monday_envio_id,
        "procedimento": "PRC-TRF4",
        "status_final": "calculado",
        "motivo_encaminhamento": (
            "Cálculo efetuado com imposto indevido a restituir"
        ),
        "requerente_nome": _get(output_payload, "requerente", "nome"),
        "cpf_cnpj": _get(output_payload, "requerente", "cpf_cnpj"),
        "requerido": demonstrativo_metadata.get("requerido"),
        "devedor": demonstrativo_metadata.get("requerido"),
        "advogado": demonstrativo_metadata.get("advogado"),
        "processo_pasta": demonstrativo.get("processo_pasta"),
        "processo_originario": identificacao.get("processo_originario")
        or demonstrativo.get("processo_pasta"),
        "processo_requisicao": identificacao.get("processo_requisicao")
        or demonstrativo.get("processo_requisicao"),
        "banco": identificacao.get("banco_comprovante") or demonstrativo.get("banco"),
        "agencia": demonstrativo.get("agencia"),
        "conta_deposito": demonstrativo.get("conta_deposito"),
        "fluxo": (output_payload.get("decisao_fluxo") or {}).get("tipo_fluxo"),
        "valor_total_pagamento": valor_total_pagamento,
        "data_levantamento": data_levantamento,
        "oc_ano": _year_from_date(data_levantamento),
        "data_preenchimento_demonstrativo": date.today().isoformat(),
        "data_deposito": _iso_date_or_original(data_levantamento),
        "mes_recebimento": dados_comprovante.get("mes_recebimento")
        or (receita_rra.get("entrada") or {}).get("mes_recebimento"),
        "numero_meses_rra": _money_to_float(
            dados_comprovante.get("numero_meses_rra")
            or (receita_rra.get("entrada") or {}).get("numero_meses")
        ),
        "imposto_retido_na_fonte": ir_retido_comprovante,
        "irrf_calculado_receita": _money_to_float(
            receita_resultado.get("irrf")
            or imposto_atualizado.get("irrf_calculado")
        ),
        "imposto_indevido": _money_to_float(
            imposto_atualizado.get("imposto_indevido")
        ),
        "selic_percentual": _money_to_float(selic.get("percentual")),
        "selic_acumulado": _money_to_float(
            imposto_atualizado.get("selic_acumulado")
        ),
        "valor_final_atualizado": imposto_retido_indevidamente_atualizado,
        "imposto_retido_indevidamente_atualizado": (
            imposto_retido_indevidamente_atualizado
        ),
        "ir_retido_monday": ir_retido_comprovante,
        "imposto_a_restituir": imposto_retido_indevidamente_atualizado,
        "json_path": str(output_path),
        "demonstrativo_pdf": demonstrativo_pdf,
        "comprovante_pdf": _first(output_payload, "arquivos_originais", index=1),
    }
    campos = {key: value for key, value in campos.items() if value not in (None, "")}
    item_name = str(monday_envio_id)
    return {
        "item_name": item_name,
        "campos": campos,
    }


def _resolve_demonstrativo_metadata(
    demonstrativo: dict[str, Any],
    demonstrativo_pdf: Any,
) -> dict[str, Any]:
    metadata = {
        "requerido": demonstrativo.get("requerido"),
        "advogado": demonstrativo.get("advogado"),
    }
    if metadata.get("requerido") and metadata.get("advogado"):
        return metadata
    if not demonstrativo_pdf:
        return metadata
    path = Path(str(demonstrativo_pdf))
    if not path.exists():
        return metadata
    campos = extract_campos_from_pdf(path)
    return {
        "requerido": metadata.get("requerido") or campos.requerido,
        "advogado": metadata.get("advogado") or campos.advogado,
    }


def _year_from_date(value: Any) -> str | None:
    iso_date = _iso_date_or_original(value)
    if not iso_date:
        return None
    match = re.search(r"\b(\d{4})-\d{2}-\d{2}\b", str(iso_date))
    if match:
        return match.group(1)
    match = re.search(r"\b\d{1,2}/\d{1,2}/(\d{4})\b", str(value or ""))
    return match.group(1) if match else None


def _iso_date_or_original(value: Any) -> str | None:
    if value in (None, ""):
        return None
    text = str(value).strip()
    iso_match = re.search(r"\b\d{4}-\d{2}-\d{2}\b", text)
    if iso_match:
        return iso_match.group(0)
    br_match = re.search(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b", text)
    if br_match:
        day, month, year = br_match.groups()
        return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"
    return text


def _read_monday_envio_state(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {
            "ultimo_envio_id": 0,
            "registros": {},
        }
    try:
        payload = _read_json(path)
    except Exception:  # noqa: BLE001
        return {
            "ultimo_envio_id": 0,
            "registros": {},
        }
    if not isinstance(payload, dict):
        return {
            "ultimo_envio_id": 0,
            "registros": {},
        }
    payload.setdefault("ultimo_envio_id", 0)
    payload.setdefault("registros", {})
    return payload


def _next_monday_envio_id(state: dict[str, Any], configured_last_id: int) -> int:
    state_last_id = int(state.get("ultimo_envio_id") or 0)
    return max(state_last_id, configured_last_id) + 1


def _write_monday_envio_state(
    path: Path,
    state: dict[str, Any],
    *,
    registro_id: str,
    monday_envio_id: int,
    item_id: int,
    item_name: str,
    output_path: Path,
) -> None:
    registros = state.setdefault("registros", {})
    registros[registro_id] = {
        "monday_envio_id": monday_envio_id,
        "item_id": item_id,
        "item_name": item_name,
        "output_path": str(output_path),
    }
    state["ultimo_envio_id"] = max(int(state.get("ultimo_envio_id") or 0), monday_envio_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(state, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _build_monday_item_name(campos: dict[str, Any]) -> str:
    requerente = campos.get("requerente_nome") or "Requerente sem nome"
    processo = (
        campos.get("processo_requisicao")
        or campos.get("processo_originario")
        or campos.get("processo_pasta")
        or "processo sem número"
    )
    valor = campos.get("valor_final_atualizado") or campos.get("imposto_indevido")
    if valor is None:
        return f"{requerente} - {processo}"
    return f"{requerente} - {processo} - {_format_money_br(valor)}"


def _format_money_br(value: Any) -> str:
    formatted = f"{float(value):,.2f}"
    return "R$ " + formatted.replace(",", "_").replace(".", ",").replace("_", ".")


def _eligibility_blockers(registro: dict[str, Any]) -> list[str]:
    status = registro.get("status", {})
    demonstrativo_pdf = registro.get("demonstrativo", {}).get("pdf_extraido_path")
    comprovante_pdf = registro.get("comprovante", {}).get("pdf_path")
    cpf_cnpj = registro.get("demonstrativo", {}).get("beneficiario_cpf_cnpj")
    numero_meses_rra = (
        registro.get("comprovante", {})
        .get("ir_retido", {})
        .get("numero_meses_rra")
    )
    blockers = []
    if len(_digits_only(str(cpf_cnpj or ""))) == 14:
        blockers.append("documento_cnpj_ignorado")
    if not numero_meses_rra:
        blockers.append("comprovante.numero_meses_rra_ausente")
    if not status.get("pronto_para_calculo"):
        blockers.append("status.pronto_para_calculo_false")
    if not status.get("tem_ir_retido"):
        blockers.append("status.tem_ir_retido_false")
    if status.get("ignorado"):
        blockers.append(f"status.ignorado:{status.get('motivo_ignorado')}")
    if not demonstrativo_pdf:
        blockers.append("demonstrativo.pdf_extraido_path_ausente")
    if not comprovante_pdf:
        blockers.append("comprovante.pdf_path_ausente")
    return blockers


def _resolve_fluxo(demonstrativo_pdf: Path, settings: Settings) -> dict[str, Any]:
    juros_mora_selic: float | None = None
    if demonstrativo_pdf.exists():
        juros_mora_selic = extract_campos_from_pdf(demonstrativo_pdf).juros_mora_selic

    if juros_mora_selic is None:
        tipo_fluxo = "juros_indeterminado_usar_planilha_por_segurança"
        prompt_path = settings.openai_calculo_juros_zerado_prompt_file
    elif abs(juros_mora_selic) < 0.005:
        tipo_fluxo = "juros_zerado_usar_planilha"
        prompt_path = settings.openai_calculo_juros_zerado_prompt_file
    else:
        tipo_fluxo = "rra_direto_juros_mora_selic_nao_zerado"
        prompt_path = settings.openai_calculo_rra_direto_prompt_file

    return {
        "tipo_fluxo": tipo_fluxo,
        "juros_mora_selic_detectado": juros_mora_selic,
        "demonstrativo_pdf_usado_na_decisao": str(demonstrativo_pdf),
        "prompt_usado": str(prompt_path),
    }


def _build_llm_processor(settings: Settings, fluxo: dict[str, Any]):
    if fluxo.get("tipo_fluxo") in {
        "juros_zerado_usar_planilha",
        "juros_indeterminado_usar_planilha_por_segurança",
    }:
        prompt = _with_spreadsheet_context(
            settings.resolved_calculo_juros_zerado_prompt(),
            settings.planilha_juros_modelo_path,
        )
    else:
        prompt = settings.resolved_calculo_rra_direto_prompt()
    return build_pdf_extractor(settings, prompt)


def _with_spreadsheet_context(prompt: str, spreadsheet_path: Path) -> str:
    if not spreadsheet_path.exists():
        logger.warning(
            "planilha_juros_modelo_not_found",
            path=str(spreadsheet_path),
        )
        return prompt
    return (
        f"{prompt}\n\n"
        "Referência textual extraída da planilha Excel modelo do escritório:\n"
        f"{_spreadsheet_reference_text(spreadsheet_path)}"
    )


def _spreadsheet_reference_text(path: Path, *, max_cells: int = 500) -> str:
    workbook = load_workbook(path, data_only=False, read_only=True)
    lines = [f"Arquivo: {path.name}"]
    cells_count = 0
    for sheet in workbook.worksheets:
        lines.append(f"Aba: {sheet.title}")
        for row in sheet.iter_rows():
            values = []
            for cell in row:
                if cell.value is None:
                    continue
                values.append(f"{cell.coordinate}={cell.value}")
            if not values:
                continue
            lines.append(" | ".join(values))
            cells_count += len(values)
            if cells_count >= max_cells:
                lines.append("Conteúdo truncado para manter o prompt compacto.")
                return "\n".join(lines)
    return "\n".join(lines)


def _record_output_id(registro: dict[str, Any]) -> str:
    demonstrativo = registro.get("demonstrativo", {})
    parts = [
        demonstrativo.get("processo_pasta"),
        demonstrativo.get("processo_requisicao"),
        demonstrativo.get("banco"),
        demonstrativo.get("conta_deposito"),
    ]
    return _safe_filename("_".join(str(part or "sem_valor") for part in parts))


def _safe_filename(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", value).strip("_") or "sem_valor"


def _total_size(paths: list[Path]) -> int:
    return sum(path.stat().st_size for path in paths)


def _optimize_attachments_for_openai(paths: list[Path], output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    optimized_paths = []
    for path in paths:
        optimized_paths.append(_optimize_pdf_for_openai(path, output_dir))
    return optimized_paths


def _optimize_pdf_for_openai(path: Path, output_dir: Path) -> Path:
    output_path = output_dir / path.name
    if path.suffix.lower() != ".pdf":
        return path

    original_size = path.stat().st_size
    doc = fitz.open(path)
    try:
        doc.save(output_path, garbage=4, deflate=True, clean=True)
    finally:
        doc.close()

    if not output_path.exists() or output_path.stat().st_size >= original_size:
        return path
    return output_path


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _calculate_receita_rra(
    parsed_data: dict[str, Any],
    *,
    timeout_seconds: int,
) -> dict[str, Any]:
    try:
        case = _get_first_case(parsed_data)
        if not case:
            return {
                "status": "nao_calculado",
                "motivo": "Caso estruturado ausente na resposta da IA",
            }
        receita = case.get("campos_para_simulador_receita") or {}
        comprovante = case.get("dados_comprovante_bancario") or {}
        month, year = _parse_month_year(receita.get("campo_4_1_mes_recebimento"))
        if month is None or year is None:
            month, year = _parse_month_year(comprovante.get("mes_recebimento"))
        if year is None and receita.get("ano_calendario") not in (None, ""):
            year = int(receita["ano_calendario"])

        missing = []
        if year is None:
            missing.append("ano extraído de campo_4_1_mes_recebimento")
        if month is None:
            missing.append("mês extraído de campo_4_1_mes_recebimento")
        if receita.get("campo_4_2_numero_meses") in (None, ""):
            missing.append("campo_4_2_numero_meses")
        if receita.get("campo_3_1_rendimentos_tributaveis") in (None, ""):
            missing.append("campo_3_1_rendimentos_tributaveis")
        if receita.get("campo_3_2_valor_dos_juros") in (None, ""):
            missing.append("campo_3_2_valor_dos_juros")
        if missing:
            return {
                "status": "nao_calculado",
                "motivo": "Campos obrigatórios ausentes",
                "campos_faltantes": missing,
            }

        rra_input = RRAInput(
            ano_calendario=int(year),
            mes_recebimento=int(month),
            numero_meses=_decimal(receita.get("campo_4_2_numero_meses")),
            rendimentos_recebidos=_decimal(
                receita.get("campo_3_1_rendimentos_tributaveis")
            ),
            juros=_decimal(receita.get("campo_3_2_valor_dos_juros")),
            previdencia_oficial=(
                _decimal(comprovante.get("deducao_previdenciaria"))
                or _decimal(receita.get("previdencia_oficial"))
            ),
        )
        result = ReceitaRRAClient(timeout_seconds=timeout_seconds).calcular(rra_input)
        return {
            "status": "calculado",
            "entrada": {
                "ano_calendario": rra_input.ano_calendario,
                "mes_recebimento": rra_input.mes_recebimento,
                "numero_meses": str(rra_input.numero_meses),
                "rendimentos_recebidos": str(rra_input.rendimentos_recebidos),
                "juros": str(rra_input.juros),
                "previdencia_oficial": str(rra_input.previdencia_oficial),
            },
            "resultado": result.to_dict(),
        }
    except Exception as exc:  # noqa: BLE001
        return {
            "status": "erro",
            "motivo": str(exc),
        }


def _apply_comprovante_deterministic_overrides(
    parsed_data: dict[str, Any],
    registro: dict[str, Any],
    demonstrativo_pdf: Path,
) -> list[dict[str, Any]]:
    ajustes = []
    ajustes.extend(_apply_comprovante_tax_overrides(parsed_data, registro))
    ajustes.extend(_apply_comprovante_month_override(parsed_data, registro))
    ajustes.extend(
        _apply_demonstrativo_receita_value_overrides(
            parsed_data,
            demonstrativo_pdf,
        )
    )
    return ajustes


def _apply_comprovante_tax_overrides(
    parsed_data: dict[str, Any],
    registro: dict[str, Any],
) -> list[dict[str, Any]]:
    ir_retido = registro.get("comprovante", {}).get("ir_retido", {})
    valor_ir_retido = _money_to_float(ir_retido.get("valor_ir_retido"))
    deducao_previdenciaria = _money_to_float(ir_retido.get("deducao_previdenciaria"))

    case = _get_first_case(parsed_data)
    if not case:
        return []

    comprovante = case.setdefault("dados_comprovante_bancario", {})
    receita = case.setdefault("campos_para_simulador_receita", {})
    previous_values = {}
    applied = {}
    numero_meses_escolhido = escolher_numero_meses_rra(
        extraido_local=ir_retido.get("numero_meses_rra"),
        extraido_ia=(
            receita.get("campo_4_2_numero_meses")
            if receita.get("campo_4_2_numero_meses") not in (None, "")
            else comprovante.get("numero_meses_rra")
        ),
    )

    if valor_ir_retido is not None:
        previous_values["dados_comprovante_bancario.imposto_retido_na_fonte"] = (
            comprovante.get("imposto_retido_na_fonte")
        )
        previous_values["dados_comprovante_bancario.tem_imposto_retido_na_fonte"] = (
            comprovante.get("tem_imposto_retido_na_fonte")
        )
        comprovante["imposto_retido_na_fonte"] = valor_ir_retido
        comprovante["tem_imposto_retido_na_fonte"] = valor_ir_retido > 0
        applied["imposto_retido_na_fonte"] = valor_ir_retido

    if deducao_previdenciaria is not None:
        previous_values["dados_comprovante_bancario.deducao_previdenciaria"] = (
            comprovante.get("deducao_previdenciaria")
        )
        previous_values["campos_para_simulador_receita.previdencia_oficial"] = (
            receita.get("previdencia_oficial")
        )
        comprovante["deducao_previdenciaria"] = deducao_previdenciaria
        receita["previdencia_oficial"] = deducao_previdenciaria
        applied["deducao_previdenciaria"] = deducao_previdenciaria

    if numero_meses_escolhido is not None:
        meses_int = int(numero_meses_escolhido)
        previous_values["dados_comprovante_bancario.numero_meses_rra"] = (
            comprovante.get("numero_meses_rra")
        )
        previous_values["campos_para_simulador_receita.campo_4_2_numero_meses"] = (
            receita.get("campo_4_2_numero_meses")
        )
        comprovante["numero_meses_rra"] = meses_int
        receita["campo_4_2_numero_meses"] = meses_int
        applied["numero_meses_rra"] = meses_int

    if not applied:
        return []

    return [
        {
            "campo": "imposto_retido_deducao_previdenciaria_numero_meses",
            "motivo": (
                "Campos fiscais do comprovante fixados deterministicamente pela "
                "extração local para evitar inversão entre IR e dedução pela IA."
            ),
            "valores_aplicados": applied,
            "valores_substituidos": previous_values,
        }
    ]


def _apply_comprovante_month_override(
    parsed_data: dict[str, Any],
    registro: dict[str, Any],
) -> list[dict[str, Any]]:
    data_levantamento = (
        registro.get("comprovante", {}).get("data_levantamento")
        or _get_first_case(parsed_data)
        .get("dados_comprovante_bancario", {})
        .get("data_levantamento")
    )
    month, year = _parse_full_date_month_year(data_levantamento)
    if month is None or year is None:
        return []

    case = _get_first_case(parsed_data)
    if not case:
        return []

    mes_recebimento = f"{month:02d}/{year:04d}"
    comprovante = case.setdefault("dados_comprovante_bancario", {})
    receita = case.setdefault("campos_para_simulador_receita", {})

    previous_values = {
        "dados_comprovante_bancario.mes_recebimento": comprovante.get(
            "mes_recebimento"
        ),
        "campos_para_simulador_receita.campo_4_1_mes_recebimento": receita.get(
            "campo_4_1_mes_recebimento"
        ),
        "campos_para_simulador_receita.ano_calendario": receita.get(
            "ano_calendario"
        ),
    }
    comprovante["mes_recebimento"] = mes_recebimento
    receita["campo_4_1_mes_recebimento"] = mes_recebimento
    receita["ano_calendario"] = year

    return [
        {
            "campo": "mes_recebimento",
            "motivo": (
                "Mês de recebimento fixado deterministicamente pela data de "
                "levantamento do comprovante bancário."
            ),
            "data_levantamento_usada": data_levantamento,
            "valor_aplicado": mes_recebimento,
            "valores_substituidos": previous_values,
        }
    ]


def _apply_demonstrativo_receita_value_overrides(
    parsed_data: dict[str, Any],
    demonstrativo_pdf: Path,
) -> list[dict[str, Any]]:
    financials = _extract_demonstrativo_receita_financials(demonstrativo_pdf)
    if not financials:
        return []

    case = _get_first_case(parsed_data)
    if not case:
        return []

    receita = case.setdefault("campos_para_simulador_receita", {})
    previous_values = {}
    applied = {}

    total_devido = financials.get("valor_total_devido_pagamento")
    if total_devido is not None and receita.get("campo_3_1_rendimentos_tributaveis") in (
        None,
        "",
    ):
        previous_values["campos_para_simulador_receita.campo_3_1_rendimentos_tributaveis"] = receita.get(
            "campo_3_1_rendimentos_tributaveis"
        )
        receita["campo_3_1_rendimentos_tributaveis"] = total_devido
        applied["campo_3_1_rendimentos_tributaveis"] = total_devido

    juros_mora = financials.get("juros_mora_pagamento")
    if juros_mora is not None and receita.get("campo_3_2_valor_dos_juros") in (
        None,
        "",
    ):
        previous_values["campos_para_simulador_receita.campo_3_2_valor_dos_juros"] = receita.get(
            "campo_3_2_valor_dos_juros"
        )
        receita["campo_3_2_valor_dos_juros"] = juros_mora
        applied["campo_3_2_valor_dos_juros"] = juros_mora

    if not applied:
        return []

    return [
        {
            "campo": "campos_receita_demonstrativo",
            "motivo": (
                "Campos do simulador RRA preenchidos deterministicamente a partir "
                "do demonstrativo para evitar cálculo bloqueado por campos nulos da IA."
            ),
            "arquivo_usado": str(demonstrativo_pdf),
            "valores_aplicados": applied,
            "valores_substituidos": previous_values,
        }
    ]


def _extract_demonstrativo_receita_financials(path: Path) -> dict[str, float]:
    if not path.exists():
        return {}

    doc = fitz.open(path)
    try:
        text = "\n".join(page.get_text() for page in doc)
    finally:
        doc.close()

    financials: dict[str, float] = {}
    total_match = re.search(
        r"Valor\s+total\s+devido\s+no\s+pagamento\s+em\s+\d{2}/\d{4}\s*"
        r"\([^)]*\)\s*:\s*([0-9.\s]+,\d{2})",
        text,
        re.IGNORECASE,
    )
    if total_match:
        financials["valor_total_devido_pagamento"] = _money_to_float(
            total_match.group(1)
        )

    final_values_match = re.search(
        r"\(\s*Principal\s+corrigido\s*:\s*([0-9.\s]+,\d{2})\s+"
        r"Juros\s+de\s+mora\s*:\s*([0-9.\s]+,\d{2})\s+"
        r"Selic\s*:\s*([0-9.\s]+,\d{2})\s*\)",
        text,
        re.IGNORECASE,
    )
    if final_values_match:
        financials["principal_corrigido_pagamento"] = _money_to_float(
            final_values_match.group(1)
        )
        financials["juros_mora_pagamento"] = _money_to_float(
            final_values_match.group(2)
        )
        financials["selic_pagamento"] = _money_to_float(final_values_match.group(3))

    return {
        key: value
        for key, value in financials.items()
        if value is not None
    }


def _calculate_imposto_indevido_atualizado(
    parsed_data: dict[str, Any],
    resultado_receita_rra: dict[str, Any],
    *,
    timeout_seconds: int,
) -> dict[str, Any]:
    try:
        if resultado_receita_rra.get("status") != "calculado":
            return {
                "status": "nao_calculado",
                "motivo": "RRA ainda não calculado",
            }

        case = _get_first_case(parsed_data)
        if not case:
            return {
                "status": "nao_calculado",
                "motivo": "Caso estruturado ausente na resposta da IA",
            }
        receita = case.get("campos_para_simulador_receita") or {}
        comprovante = case.get("dados_comprovante_bancario") or {}
        imposto_retido = comprovante.get("imposto_retido_na_fonte")
        irrf_calculado = (
            (resultado_receita_rra.get("resultado") or {}).get("irrf")
        )
        month, year = _parse_month_year(receita.get("campo_4_1_mes_recebimento"))
        if month is None or year is None:
            month, year = _parse_month_year(comprovante.get("mes_recebimento"))

        missing = []
        if imposto_retido in (None, ""):
            missing.append("dados_comprovante_bancario.imposto_retido_na_fonte")
        if irrf_calculado in (None, ""):
            missing.append("resultado_receita_rra.resultado.irrf")
        if month is None or year is None:
            missing.append("campo_4_1_mes_recebimento")
        if missing:
            return {
                "status": "nao_calculado",
                "motivo": "Campos obrigatórios ausentes",
                "campos_faltantes": missing,
            }

        selic = SicalcSelicClient(timeout_seconds=timeout_seconds).consultar_acumulada(
            vencimento_mes=month,
            vencimento_ano=year,
        )
        return calcular_imposto_indevido_atualizado(
            imposto_retido=_decimal(imposto_retido),
            irrf_calculado=_decimal(irrf_calculado),
            selic=selic,
        )
    except Exception as exc:  # noqa: BLE001
        return {
            "status": "erro",
            "motivo": str(exc),
        }


def _normalize_parsed_extraction(parsed_data: dict[str, Any]) -> dict[str, Any]:
    """Garante estrutura com `casos` mesmo se a LLM devolveu JSON truncado."""
    if not isinstance(parsed_data, dict):
        return {}
    if parsed_data.get("casos"):
        return parsed_data
    bruto = parsed_data.get("conteudo_bruto")
    if isinstance(bruto, str) and bruto.strip():
        recovered = parse_llm_json(bruto)
        if isinstance(recovered, dict) and recovered.get("casos"):
            return recovered
    return parsed_data


def _get_first_case(parsed_data: dict[str, Any]) -> dict[str, Any]:
    cases = parsed_data.get("casos") or []
    if not cases:
        return {}
    return cases[0]


def _get(payload: dict[str, Any], *keys: str) -> Any:
    value: Any = payload
    for key in keys:
        if not isinstance(value, dict):
            return None
        value = value.get(key)
    return value


def _first(payload: dict[str, Any], *keys: str, index: int = 0) -> Any:
    value = _get(payload, *keys)
    if isinstance(value, list) and len(value) > index:
        return value[index]
    return None


def _parse_month_year(value: Any) -> tuple[int | None, int | None]:
    if value is None:
        return None, None
    text = str(value).strip()
    match = re.search(r"(?P<month>\d{1,2})\s*/\s*(?P<year>\d{4})", text)
    if match:
        return int(match.group("month")), int(match.group("year"))

    month_names = {
        "janeiro": 1,
        "fevereiro": 2,
        "março": 3,
        "marco": 3,
        "abril": 4,
        "maio": 5,
        "junho": 6,
        "julho": 7,
        "agosto": 8,
        "setembro": 9,
        "outubro": 10,
        "novembro": 11,
        "dezembro": 12,
    }
    lowered = text.lower()
    month = month_names.get(lowered)
    if month:
        return month, None
    if text.isdigit() and 1 <= int(text) <= 12:
        return int(text), None
    return None, None


def _parse_full_date_month_year(value: Any) -> tuple[int | None, int | None]:
    if value is None:
        return None, None
    text = str(value).strip()
    match = re.search(r"\b\d{1,2}\s*/\s*(?P<month>\d{1,2})\s*/\s*(?P<year>\d{4})\b", text)
    if match:
        return int(match.group("month")), int(match.group("year"))
    match = re.search(r"\b(?P<year>\d{4})-(?P<month>\d{1,2})-\d{1,2}\b", text)
    if match:
        return int(match.group("month")), int(match.group("year"))
    return None, None


def _decimal(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    text = str(value).strip().replace("R$", "").replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        return Decimal("0")


def _money_to_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    return float(_decimal(value))


def _digits_only(value: str) -> str:
    return re.sub(r"\D+", "", str(value or ""))
