from __future__ import annotations

import json
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import structlog
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from recupere_monday.config import Settings
from recupere_monday.storage.operational_db import OperationalDBRepository

logger = structlog.get_logger(__name__)


@dataclass
class RequerentePayloadsResult:
    total_registros: int = 0
    total_requerentes: int = 0
    total_requerentes_aptos: int = 0
    total_requerentes_inaptos: int = 0
    registros_aptos: int = 0
    registros_inaptos_sem_ir: int = 0
    registros_inaptos_sem_numero_meses: int = 0
    registros_inaptos_falha_comprovante: int = 0
    registros_ignorados_cnpj: int = 0
    arquivos_gerados: int = 0
    arquivos_inaptos_gerados: int = 0
    arquivos_sem_numero_meses_gerados: int = 0
    arquivos_falha_comprovante_gerados: int = 0
    arquivos_cnpj_gerados: int = 0
    output_dir: Path | None = None
    inaptos_output_dir: Path | None = None
    controle_path: Path | None = None
    failed: int = 0
    errors: list[str] = field(default_factory=list)


class RequerentePayloadsOrchestrator:
    """Gera a base individualizada por requerente para cálculo e controle."""

    def __init__(self, settings: Settings) -> None:
        self._demonstrativos_dir = settings.demonstrativo_campos_output_path
        self._demonstrativos_pdf_dir = settings.extracted_demonstrativos_path
        self._status_dirs = {
            "CEF": settings.comprovantes_caixa_status_output_path,
            "BB": settings.comprovantes_bb_status_output_path,
        }
        self._ir_report_path = settings.comprovantes_ir_retido_output_path
        self._output_dir = settings.requerentes_output_path
        self._inaptos_output_dir = settings.requerentes_inaptos_output_path
        self._sem_ir_output_dir = self._inaptos_output_dir / "sem_ir"
        self._sem_numero_meses_output_dir = (
            self._inaptos_output_dir / "sem_numero_meses"
        )
        self._falha_comprovante_output_dir = (
            self._inaptos_output_dir / "falha_comprovante"
        )
        self._cnpj_output_dir = self._inaptos_output_dir / "cnpj"
        self._controle_path = settings.requerentes_controle_output_path
        self._prompt_path = settings.openai_demonstrativo_prompt_file
        self._operational_db = OperationalDBRepository.from_settings(settings)

    def run(
        self,
        *,
        process_filter: str | None = None,
        clear_previous: bool = True,
    ) -> RequerentePayloadsResult:
        result = RequerentePayloadsResult(
            output_dir=self._output_dir,
            inaptos_output_dir=self._inaptos_output_dir,
            controle_path=self._controle_path,
        )
        if not self._demonstrativos_dir.exists():
            result.failed += 1
            result.errors.append(
                f"Diretório de demonstrativos não encontrado: {self._demonstrativos_dir}"
            )
            return result

        ir_by_pdf_path = self._load_ir_by_pdf_path()
        demonstrativo_files = sorted(
            path
            for path in self._demonstrativos_dir.rglob("*.json")
            if process_filter is None or process_filter in path.parts
        )
        result.total_registros = len(demonstrativo_files)

        registros = []
        for demonstrativo_path in demonstrativo_files:
            try:
                registros.append(self._build_registro(demonstrativo_path, ir_by_pdf_path))
            except Exception as exc:  # noqa: BLE001
                result.failed += 1
                result.errors.append(f"{demonstrativo_path.name}: {exc}")
                logger.exception(
                    "requerente_registro_failed",
                    demonstrativo=str(demonstrativo_path),
                    error=str(exc),
                )

        grupos = self._group_by_requerente(registros)
        result.total_requerentes = len(grupos)
        self._output_dir.mkdir(parents=True, exist_ok=True)
        self._sem_ir_output_dir.mkdir(parents=True, exist_ok=True)
        self._sem_numero_meses_output_dir.mkdir(parents=True, exist_ok=True)
        self._falha_comprovante_output_dir.mkdir(parents=True, exist_ok=True)
        self._cnpj_output_dir.mkdir(parents=True, exist_ok=True)
        if clear_previous:
            self._clear_previous_payloads()

        controle_rows: list[dict[str, Any]] = []
        for requerente_key, group_records in sorted(grupos.items()):
            aptos = [
                item
                for item in group_records
                if item["status"]["tem_ir_retido"]
                and item["status"]["tem_numero_meses_rra"]
                and not item["status"]["ignorado"]
            ]
            sem_ir = [
                item
                for item in group_records
                if item["status"]["comprovante_encontrado"]
                and not item["status"]["tem_ir_retido"]
                and not item["status"]["ignorado"]
            ]
            cnpj_ignorados = [
                item
                for item in group_records
                if item["status"]["ignorado"]
                and item["status"]["motivo_ignorado"] == "cnpj_ignorado"
            ]
            sem_numero_meses = [
                item
                for item in group_records
                if item["status"]["comprovante_encontrado"]
                and item["status"]["tem_ir_retido"]
                and not item["status"]["tem_numero_meses_rra"]
                and not item["status"]["ignorado"]
            ]
            falha_comprovante = [
                item
                for item in group_records
                if item["status"]["falha_comprovante"]
                and not item["status"]["ignorado"]
            ]

            if aptos:
                payload = self._build_payload(
                    requerente_key,
                    aptos,
                    tipo_payload="calculo_ir_requerente",
                )
                payload_path = self._write_case_payload(
                    self._output_dir / requerente_key,
                    payload,
                    aptos,
                )
                self._persist_payload(
                    payload,
                    payload_path,
                    classificacao="apto_com_ir",
                )
                result.arquivos_gerados += 1
                result.total_requerentes_aptos += 1
                result.registros_aptos += len(aptos)
                for registro in aptos:
                    controle_rows.append(
                        self._build_controle_row(
                            payload,
                            registro,
                            payload_path,
                            classificacao="apto_com_ir",
                        )
                    )

            if sem_ir:
                payload_inapto = self._build_payload(
                    requerente_key,
                    sem_ir,
                    tipo_payload="casos_inaptos_sem_ir",
                )
                payload_inapto["motivo_inaptidao"] = (
                    "Registros sem IR retido no comprovante bancário."
                )
                payload_inapto_path = self._write_case_payload(
                    self._sem_ir_output_dir / requerente_key,
                    payload_inapto,
                    sem_ir,
                )
                self._persist_payload(
                    payload_inapto,
                    payload_inapto_path,
                    classificacao="inapto_sem_ir",
                )
                result.arquivos_inaptos_gerados += 1
                result.total_requerentes_inaptos += 1
                result.registros_inaptos_sem_ir += len(sem_ir)
                for registro in sem_ir:
                    controle_rows.append(
                        self._build_controle_row(
                            payload_inapto,
                            registro,
                            payload_inapto_path,
                            classificacao="inapto_sem_ir",
                        )
                    )

            if sem_numero_meses:
                payload_sem_numero_meses = self._build_payload(
                    requerente_key,
                    sem_numero_meses,
                    tipo_payload="casos_inaptos_sem_numero_meses",
                )
                payload_sem_numero_meses["motivo_inaptidao"] = (
                    "Registros com IR retido, mas sem Número de Meses RRA no comprovante bancário."
                )
                payload_sem_numero_meses_path = self._write_case_payload(
                    self._sem_numero_meses_output_dir / requerente_key,
                    payload_sem_numero_meses,
                    sem_numero_meses,
                )
                self._persist_payload(
                    payload_sem_numero_meses,
                    payload_sem_numero_meses_path,
                    classificacao="inapto_sem_numero_meses",
                )
                result.arquivos_sem_numero_meses_gerados += 1
                result.total_requerentes_inaptos += 1
                result.registros_inaptos_sem_numero_meses += len(sem_numero_meses)
                for registro in sem_numero_meses:
                    controle_rows.append(
                        self._build_controle_row(
                            payload_sem_numero_meses,
                            registro,
                            payload_sem_numero_meses_path,
                            classificacao="inapto_sem_numero_meses",
                        )
                    )

            if falha_comprovante:
                payload_falha_comprovante = self._build_payload(
                    requerente_key,
                    falha_comprovante,
                    tipo_payload="casos_inaptos_falha_comprovante",
                )
                payload_falha_comprovante["motivo_inaptidao"] = (
                    "Registros sem comprovante bancário válido por falha, erro ou não localização na coleta."
                )
                payload_falha_comprovante_path = self._write_case_payload(
                    self._falha_comprovante_output_dir / requerente_key,
                    payload_falha_comprovante,
                    falha_comprovante,
                )
                self._persist_payload(
                    payload_falha_comprovante,
                    payload_falha_comprovante_path,
                    classificacao="inapto_falha_comprovante",
                )
                result.arquivos_falha_comprovante_gerados += 1
                result.total_requerentes_inaptos += 1
                result.registros_inaptos_falha_comprovante += len(falha_comprovante)
                for registro in falha_comprovante:
                    controle_rows.append(
                        self._build_controle_row(
                            payload_falha_comprovante,
                            registro,
                            payload_falha_comprovante_path,
                            classificacao="inapto_falha_comprovante",
                        )
                    )

            if cnpj_ignorados:
                payload_cnpj = self._build_payload(
                    requerente_key,
                    cnpj_ignorados,
                    tipo_payload="casos_ignorados_cnpj",
                )
                payload_cnpj["motivo_inaptidao"] = (
                    "CNPJ ignorado por regra de negócio. O fluxo processa apenas pessoas físicas."
                )
                payload_cnpj_path = self._write_case_payload(
                    self._cnpj_output_dir / requerente_key,
                    payload_cnpj,
                    cnpj_ignorados,
                )
                self._persist_payload(
                    payload_cnpj,
                    payload_cnpj_path,
                    classificacao="ignorado_cnpj",
                )
                result.arquivos_cnpj_gerados += 1
                result.registros_ignorados_cnpj += len(cnpj_ignorados)
                for registro in cnpj_ignorados:
                    controle_rows.append(
                        self._build_controle_row(
                            payload_cnpj,
                            registro,
                            payload_cnpj_path,
                            classificacao="ignorado_cnpj",
                        )
                    )

        self._write_excel(controle_rows)
        logger.info(
            "requerente_payloads_finished",
            total_registros=result.total_registros,
            total_requerentes=result.total_requerentes,
            total_requerentes_aptos=result.total_requerentes_aptos,
            total_requerentes_inaptos=result.total_requerentes_inaptos,
            registros_aptos=result.registros_aptos,
            registros_inaptos_sem_ir=result.registros_inaptos_sem_ir,
            registros_inaptos_sem_numero_meses=(
                result.registros_inaptos_sem_numero_meses
            ),
            registros_inaptos_falha_comprovante=(
                result.registros_inaptos_falha_comprovante
            ),
            registros_ignorados_cnpj=result.registros_ignorados_cnpj,
            arquivos_gerados=result.arquivos_gerados,
            arquivos_inaptos_gerados=result.arquivos_inaptos_gerados,
            arquivos_sem_numero_meses_gerados=(
                result.arquivos_sem_numero_meses_gerados
            ),
            arquivos_falha_comprovante_gerados=(
                result.arquivos_falha_comprovante_gerados
            ),
            arquivos_cnpj_gerados=result.arquivos_cnpj_gerados,
            failed=result.failed,
            output_dir=str(self._output_dir),
            inaptos_output_dir=str(self._inaptos_output_dir),
            controle=str(self._controle_path),
        )
        return result

    def _build_registro(
        self,
        demonstrativo_path: Path,
        ir_by_pdf_path: dict[str, dict[str, Any]],
    ) -> dict[str, Any]:
        demonstrativo = _read_json(demonstrativo_path)
        relative = demonstrativo_path.relative_to(self._demonstrativos_dir)
        status_path = self._status_path(demonstrativo, relative)
        status_payload = _read_json(status_path) if status_path and status_path.exists() else {}
        status = status_payload.get("status", {})
        pdf_path = status.get("pdf_path")
        ir_data = ir_by_pdf_path.get(_normalize_path(pdf_path)) if pdf_path else None
        banco = str(demonstrativo.get("banco", "")).upper()
        ignorado = bool(status.get("skipped")) and status.get("reason") == "cnpj_ignorado"
        comprovante_encontrado = bool(status.get("success")) and bool(pdf_path)
        falha_comprovante = (
            status_path is not None
            and status_path.exists()
            and not comprovante_encontrado
            and status.get("reason") != "cnpj_ignorado"
        )
        tem_ir_retido = bool(ir_data.get("tem_ir_retido")) if ir_data else False
        numero_meses_rra = _clean_text((ir_data or {}).get("numero_meses_rra"))
        tem_numero_meses_rra = bool(numero_meses_rra)
        beneficiario_nome = _clean_text(demonstrativo.get("beneficiario_nome")) or _clean_text(
            (ir_data or {}).get("requerente_nome")
        )
        beneficiario_cpf_cnpj = _clean_text(
            demonstrativo.get("beneficiario_cpf_cnpj")
        ) or _clean_text((ir_data or {}).get("requerente_cpf_cnpj"))
        document_digits = _digits_only(beneficiario_cpf_cnpj or "")
        is_cnpj = len(document_digits) == 14
        ignorado = ignorado or is_cnpj
        motivo_ignorado = "cnpj_ignorado" if is_cnpj else status.get("reason")

        return {
            "id": _build_record_id(demonstrativo, relative),
            "status": {
                "comprovante_encontrado": comprovante_encontrado,
                "tem_ir_retido": tem_ir_retido,
                "tem_numero_meses_rra": tem_numero_meses_rra,
                "falha_comprovante": falha_comprovante,
                "ignorado": ignorado,
                "motivo_ignorado": motivo_ignorado if ignorado else None,
                "erro_comprovante": status.get("error"),
                "motivo_falha_comprovante": (
                    _resolve_comprovante_failure_reason(status)
                    if falha_comprovante
                    else None
                ),
                "pronto_para_calculo": (
                    comprovante_encontrado
                    and tem_ir_retido
                    and tem_numero_meses_rra
                    and not falha_comprovante
                    and not ignorado
                ),
            },
            "demonstrativo": {
                "arquivo": demonstrativo.get("arquivo"),
                "processo_pasta": demonstrativo.get("processo_pasta"),
                "processo_requisicao": demonstrativo.get("processo_requisicao"),
                "beneficiario_nome": beneficiario_nome,
                "beneficiario_cpf_cnpj": beneficiario_cpf_cnpj,
                "requerido": demonstrativo.get("requerido"),
                "advogado": demonstrativo.get("advogado"),
                "banco": banco,
                "agencia": demonstrativo.get("agencia"),
                "conta_deposito": demonstrativo.get("conta_deposito"),
                "mes_pagamento": demonstrativo.get("mes_pagamento"),
                "data_disponibilidade_saque": demonstrativo.get(
                    "data_disponibilidade_saque"
                ),
                "json_path": str(demonstrativo_path),
                "pdf_extraido_path": str(self._demonstrativo_pdf_path(relative)),
            },
            "comprovante": self._build_comprovante(status_path, status, ir_data),
        }

    def _build_comprovante(
        self,
        status_path: Path | None,
        status: dict[str, Any],
        ir_data: dict[str, Any] | None,
    ) -> dict[str, Any]:
        comprovante_data = status.get("comprovante_data") or {}
        account = status.get("account") or {}
        return {
            "status_json_path": str(status_path) if status_path else None,
            "pdf_path": status.get("pdf_path"),
            "success": status.get("success"),
            "skipped": status.get("skipped"),
            "reason": status.get("reason"),
            "message": status.get("message"),
            "error": status.get("error"),
            "attempts": status.get("attempts"),
            "account": account,
            "requerente_nome": ir_data.get("requerente_nome") if ir_data else None,
            "requerente_cpf_cnpj": ir_data.get("requerente_cpf_cnpj") if ir_data else None,
            "data_levantamento": comprovante_data.get("data_levantamento"),
            "valor_bruto_resgate": comprovante_data.get("valor_bruto_resgate"),
            "valor_liquido_resgate": comprovante_data.get("valor_liquido_resgate"),
            "periodo_usado": comprovante_data.get("periodo_usado"),
            "periodo_estrategia": comprovante_data.get("periodo_estrategia"),
            "ir_retido": {
                "tem_ir_retido": bool(ir_data.get("tem_ir_retido")) if ir_data else False,
                "valor_ir_retido": ir_data.get("valor_ir_retido") if ir_data else None,
                "campo_origem": ir_data.get("campo_origem") if ir_data else None,
                "deducao_previdenciaria": (
                    ir_data.get("deducao_previdenciaria") if ir_data else None
                ),
                "numero_meses_rra": ir_data.get("numero_meses_rra") if ir_data else None,
            },
        }

    def _build_payload(
        self,
        requerente_key: str,
        registros: list[dict[str, Any]],
        *,
        tipo_payload: str,
    ) -> dict[str, Any]:
        requerente = self._resolve_requerente(registros)
        prontos = [item for item in registros if item["status"]["pronto_para_calculo"]]
        com_ir = [item for item in registros if item["status"]["tem_ir_retido"]]
        sem_ir = [
            item
            for item in registros
            if not item["status"]["tem_ir_retido"] and not item["status"]["ignorado"]
        ]
        ignorados = [item for item in registros if item["status"]["ignorado"]]

        instrucoes = _build_instrucoes(tipo_payload)
        return {
            "tipo_payload": tipo_payload,
            "prompt_referencia": str(self._prompt_path),
            "requerente_key": requerente_key,
            "requerente": requerente,
            "summary": {
                "total_registros": len(registros),
                "registros_prontos_para_calculo": len(prontos),
                "registros_com_ir_retido": len(com_ir),
                "registros_sem_ir_retido": len(sem_ir),
                "registros_ignorados": len(ignorados),
                "tem_ir_retido": bool(com_ir),
                "pronto_para_requisicao": bool(prontos),
            },
            "instrucoes_para_ia": instrucoes,
            "registros": registros,
        }

    @staticmethod
    def _write_payload(output_dir: Path, payload: dict[str, Any]) -> Path:
        output_dir.mkdir(parents=True, exist_ok=True)
        payload_path = output_dir / "payload_calculo.json"
        payload_path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return payload_path

    def _write_case_payload(
        self,
        output_dir: Path,
        payload: dict[str, Any],
        registros: list[dict[str, Any]],
    ) -> Path:
        payload_path = self._write_payload(output_dir, payload)
        self._materialize_related_files(output_dir, registros)
        return payload_path

    def _materialize_related_files(
        self,
        case_dir: Path,
        registros: list[dict[str, Any]],
    ) -> None:
        """Copia demonstrativo, comprovante e JSONs relacionados para a pasta do caso."""
        case_dir.mkdir(parents=True, exist_ok=True)
        for index, registro in enumerate(registros, start=1):
            suffix = "" if len(registros) == 1 else f"_{index}"
            demonstrativo = registro.get("demonstrativo") or {}
            comprovante = registro.get("comprovante") or {}

            _copy_related_file(
                demonstrativo.get("pdf_extraido_path"),
                case_dir / f"demonstrativo{suffix}.pdf",
            )
            _copy_related_file(
                demonstrativo.get("json_path"),
                case_dir / f"demonstrativo_campos{suffix}.json",
            )
            _copy_related_file(
                comprovante.get("pdf_path"),
                case_dir / f"comprovante{suffix}.pdf",
            )
            _copy_related_file(
                comprovante.get("status_json_path"),
                case_dir / f"comprovante_status{suffix}.json",
            )

    def _persist_payload(
        self,
        payload: dict[str, Any],
        payload_path: Path,
        *,
        classificacao: str,
    ) -> None:
        if not self._operational_db:
            return
        self._operational_db.persist_requerente_payload(
            payload=payload,
            payload_path=payload_path,
            classificacao=classificacao,
        )

    @staticmethod
    def _resolve_requerente(registros: list[dict[str, Any]]) -> dict[str, Any]:
        for registro in registros:
            demonstrativo = registro["demonstrativo"]
            nome = _clean_text(demonstrativo.get("beneficiario_nome"))
            cpf_cnpj = _clean_text(demonstrativo.get("beneficiario_cpf_cnpj"))
            if nome or cpf_cnpj:
                return {"nome": nome, "cpf_cnpj": cpf_cnpj}
        return {"nome": None, "cpf_cnpj": None}

    def _group_by_requerente(
        self,
        registros: list[dict[str, Any]],
    ) -> dict[str, list[dict[str, Any]]]:
        groups: dict[str, list[dict[str, Any]]] = {}
        for registro in registros:
            demonstrativo = registro["demonstrativo"]
            group_key = _build_case_folder_name(
                demonstrativo.get("beneficiario_nome"),
                demonstrativo.get("beneficiario_cpf_cnpj"),
                demonstrativo.get("processo_pasta"),
            )
            groups.setdefault(group_key, []).append(registro)
        return groups

    def _status_path(self, demonstrativo: dict[str, Any], relative: Path) -> Path | None:
        banco = str(demonstrativo.get("banco", "")).upper()
        status_dir = self._status_dirs.get(banco)
        if not status_dir:
            return None
        return status_dir / relative

    def _demonstrativo_pdf_path(self, relative_json_path: Path) -> Path:
        return self._demonstrativos_pdf_dir / relative_json_path.with_suffix(".pdf")

    def _load_ir_by_pdf_path(self) -> dict[str, dict[str, Any]]:
        if not self._ir_report_path.exists():
            return {}
        payload = _read_json(self._ir_report_path)
        return {
            _normalize_path(record.get("pdf_path")): record
            for record in payload.get("todos", [])
            if record.get("pdf_path")
        }

    @staticmethod
    def _build_controle_row(
        payload: dict[str, Any],
        registro: dict[str, Any],
        payload_path: Path,
        *,
        classificacao: str,
    ) -> dict[str, Any]:
        demonstrativo = registro["demonstrativo"]
        comprovante = registro["comprovante"]
        status = registro["status"]
        ir = comprovante["ir_retido"]
        return {
            "requerente_key": payload["requerente_key"],
            "classificacao": classificacao,
            "requerente_nome": payload["requerente"]["nome"],
            "cpf_cnpj": payload["requerente"]["cpf_cnpj"],
            "processo_pasta": demonstrativo.get("processo_pasta"),
            "processo_requisicao": demonstrativo.get("processo_requisicao"),
            "banco": demonstrativo.get("banco"),
            "agencia": demonstrativo.get("agencia"),
            "conta_deposito": demonstrativo.get("conta_deposito"),
            "comprovante_encontrado": status["comprovante_encontrado"],
            "tem_ir_retido": status["tem_ir_retido"],
            "tem_numero_meses_rra": status["tem_numero_meses_rra"],
            "falha_comprovante": status["falha_comprovante"],
            "valor_ir_retido": ir.get("valor_ir_retido"),
            "numero_meses_rra": ir.get("numero_meses_rra"),
            "data_levantamento": comprovante.get("data_levantamento"),
            "pronto_para_calculo": status["pronto_para_calculo"],
            "ignorado": status["ignorado"],
            "motivo_ignorado": status["motivo_ignorado"],
            "erro_comprovante": status["erro_comprovante"],
            "motivo_falha_comprovante": status["motivo_falha_comprovante"],
            "payload_calculo_path": str(payload_path),
            "demonstrativo_pdf_path": demonstrativo.get("pdf_extraido_path"),
            "comprovante_pdf_path": comprovante.get("pdf_path"),
        }

    def _write_excel(self, rows: list[dict[str, Any]]) -> None:
        self._controle_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Controle"

        headers = [
            "requerente_key",
            "classificacao",
            "requerente_nome",
            "cpf_cnpj",
            "processo_pasta",
            "processo_requisicao",
            "banco",
            "agencia",
            "conta_deposito",
            "comprovante_encontrado",
            "tem_ir_retido",
            "tem_numero_meses_rra",
            "falha_comprovante",
            "valor_ir_retido",
            "numero_meses_rra",
            "data_levantamento",
            "pronto_para_calculo",
            "ignorado",
            "motivo_ignorado",
            "erro_comprovante",
            "motivo_falha_comprovante",
            "payload_calculo_path",
            "demonstrativo_pdf_path",
            "comprovante_pdf_path",
        ]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")

        for row in rows:
            sheet.append([row.get(header) for header in headers])

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_index, header in enumerate(headers, start=1):
            max_length = max(
                len(str(sheet.cell(row=row_index, column=column_index).value or ""))
                for row_index in range(1, sheet.max_row + 1)
            )
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(max_length + 2, len(header) + 2),
                60,
            )

        workbook.save(self._controle_path)

    def _clear_previous_payloads(self) -> None:
        for output_dir in (self._output_dir, self._inaptos_output_dir):
            if output_dir.exists():
                shutil.rmtree(output_dir)


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _build_record_id(demonstrativo: dict[str, Any], relative: Path) -> str:
    processo_pasta = demonstrativo.get("processo_pasta") or relative.parent.name
    processo = demonstrativo.get("processo_requisicao") or relative.stem
    banco = demonstrativo.get("banco") or "sem_banco"
    conta = demonstrativo.get("conta_deposito") or "sem_conta"
    return f"{processo_pasta}|{processo}|{banco}|{conta}"


def _build_case_folder_name(
    nome: str | None,
    cpf_cnpj: str | None,
    processo_pasta: str | None,
) -> str:
    """Pasta unificada: '{cpf} - {processo_pasta}'."""
    digits = _digits_only(cpf_cnpj or "")
    if digits:
        identity = digits
    elif nome:
        identity = _safe_folder_segment(nome)
    else:
        identity = "sem_identificacao"
    processo = _safe_folder_segment(processo_pasta or "") or "sem_processo"
    return f"{identity} - {processo}"


def _build_group_key(nome: str | None, cpf_cnpj: str | None) -> str:
    if cpf_cnpj:
        return _digits_only(cpf_cnpj)
    if nome:
        return _safe_filename(nome).lower()
    return "requerente_sem_identificacao"


def _copy_related_file(source: Any, destination: Path) -> bool:
    if not source:
        return False
    src = Path(str(source))
    if not src.exists() or not src.is_file():
        return False
    try:
        if src.resolve() == destination.resolve():
            return True
    except OSError:
        pass
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, destination)
    return True


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    for marker in ("alvará).", "Federal."):
        if marker in text:
            text = text.split(marker)[-1].strip()
    return text or None


def _resolve_comprovante_failure_reason(status: dict[str, Any]) -> str | None:
    for key in ("reason", "error", "message"):
        value = _clean_text(status.get(key))
        if value:
            return value
    if status.get("success") is False:
        return "status_sem_sucesso"
    return None


def _normalize_path(value: Any) -> str:
    return str(value or "").replace("\\", "/").lower()


def _digits_only(value: str) -> str:
    return re.sub(r"\D+", "", str(value or ""))


def _safe_filename(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", value).strip("_") or "sem_valor"


def _safe_folder_segment(value: str) -> str:
    """Remove só caracteres inválidos no Windows; mantém espaços e hífen."""
    text = re.sub(r'[<>:"/\\|?*]+', "_", str(value or ""))
    text = re.sub(r"\s+", " ", text).strip(" .")
    return text or "sem_valor"


def _build_instrucoes(tipo_payload: str) -> dict[str, str]:
    if tipo_payload == "casos_inaptos_falha_comprovante":
        return {
            "objetivo": (
                "Auditar registros cujo comprovante bancário não foi coletado, "
                "não foi localizado ou retornou erro na coleta."
            ),
            "regra": (
                "Use este payload para resolver a pendência de coleta. Estes casos "
                "não devem ser enviados para cálculo pela IA até existir comprovante válido."
            ),
        }
    if tipo_payload == "casos_inaptos_sem_numero_meses":
        return {
            "objetivo": (
                "Auditar registros com IR retido, mas sem Número de Meses RRA "
                "identificado no comprovante bancário. Estes casos não devem ser "
                "enviados para cálculo pela IA."
            ),
            "regra": (
                "Use este payload apenas como trilha de auditoria dos casos que "
                "dependem de correção manual do número de meses RRA."
            ),
        }
    if tipo_payload == "casos_ignorados_cnpj":
        return {
            "objetivo": (
                "Auditar registros de pessoa jurídica ignorados pelo fluxo. "
                "Estes casos não devem ser enviados para cálculo pela IA."
            ),
            "regra": (
                "Use este payload apenas como trilha de auditoria dos casos CNPJ, "
                "pois o processamento é restrito a pessoas físicas."
            ),
        }
    if tipo_payload == "casos_inaptos_sem_ir":
        return {
            "objetivo": (
                "Auditar registros sem IR retido no comprovante bancário. "
                "Estes casos não devem ser enviados para cálculo pela IA."
            ),
            "regra": (
                "Use este payload apenas como trilha de auditoria dos casos separados "
                "por ausência de IR retido."
            ),
        }
    return {
        "objetivo": (
            "Calcular os valores de restituição/IR para este requerente usando "
            "os dados do demonstrativo inicial e do comprovante de levantamento."
        ),
        "regra": (
            "Use apenas os dados deste JSON. Este arquivo representa um único "
            "requerente e deve ser enviado individualmente ao modelo."
        ),
    }
