from __future__ import annotations

import json
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import structlog
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from recupere_monday.config import Settings

logger = structlog.get_logger(__name__)


@dataclass
class CalculosExcelExportResult:
    total_arquivos: int = 0
    linhas_exportadas: int = 0
    calculados: int = 0
    nao_calculados: int = 0
    output_path: Path | None = None
    failed: int = 0
    errors: list[str] = field(default_factory=list)


class CalculosExcelExporter:
    """Exporta os JSONs finais de cálculo para uma planilha executiva."""

    def __init__(self, settings: Settings) -> None:
        self._input_dir = settings.requerentes_output_path
        self._output_path = settings.calculos_excel_output_path

    def run(self) -> CalculosExcelExportResult:
        result = CalculosExcelExportResult(output_path=self._output_path)
        rows: list[dict[str, Any]] = []

        for json_path in self._calculo_files():
            result.total_arquivos += 1
            try:
                payload = _read_json(json_path)
                rows.append(_build_result_row(payload, json_path))
            except Exception as exc:  # noqa: BLE001
                result.failed += 1
                result.errors.append(f"{json_path}: {exc}")
                logger.exception(
                    "calculos_excel_export_record_failed",
                    path=str(json_path),
                    error=str(exc),
                )

        self._write_workbook(rows)
        result.linhas_exportadas = len(rows)
        result.calculados = sum(1 for row in rows if row["status_final"] == "calculado")
        result.nao_calculados = len(rows) - result.calculados
        logger.info(
            "calculos_excel_export_finished",
            total_arquivos=result.total_arquivos,
            linhas_exportadas=result.linhas_exportadas,
            calculados=result.calculados,
            nao_calculados=result.nao_calculados,
            failed=result.failed,
            output=str(self._output_path),
        )
        return result

    def _calculo_files(self) -> list[Path]:
        if not self._input_dir.exists():
            return []
        return sorted(
            path
            for path in self._input_dir.rglob("calculos_ia/*.json")
            if not path.name.startswith("_")
        )

    def _write_workbook(self, rows: list[dict[str, Any]]) -> None:
        self._output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        resumo = workbook.active
        resumo.title = "Resumo"
        resultados = workbook.create_sheet("Resultados")

        _write_summary_sheet(resumo, rows)
        _write_results_sheet(resultados, rows)
        workbook.save(self._output_path)


def _build_result_row(payload: dict[str, Any], json_path: Path) -> dict[str, Any]:
    caso = _first(payload, "resultado", "casos") or {}
    identificacao = caso.get("identificacao") or {}
    dados_demo = caso.get("dados_demonstrativo") or {}
    dados_comprovante = caso.get("dados_comprovante_bancario") or {}
    simulador = caso.get("campos_para_simulador_receita") or {}
    receita_rra = payload.get("resultado_receita_rra") or {}
    receita_resultado = receita_rra.get("resultado") or {}
    imposto_atualizado = (
        payload.get("resultado_imposto_retido_indevidamente_atualizado") or {}
    )
    selic = imposto_atualizado.get("selic") or {}
    decisao_fluxo = payload.get("decisao_fluxo") or {}

    rra_status = receita_rra.get("status")
    atualizado_status = imposto_atualizado.get("status")
    status_final = (
        "calculado"
        if rra_status == "calculado" and atualizado_status == "calculado"
        else "nao_calculado"
    )
    alertas = []
    alertas.extend(caso.get("alertas") or [])
    alertas.extend((payload.get("resultado") or {}).get("alertas_gerais") or [])

    return {
        "status_final": status_final,
        "rra_status": rra_status,
        "imposto_atualizado_status": atualizado_status,
        "motivo_nao_calculado": _join_values(
            receita_rra.get("motivo"),
            imposto_atualizado.get("motivo"),
            receita_rra.get("campos_faltantes"),
        ),
        "requerente_nome": _get(payload, "requerente", "nome"),
        "cpf_cnpj": _get(payload, "requerente", "cpf_cnpj"),
        "processo_originario": identificacao.get("processo_originario"),
        "processo_requisicao": identificacao.get("processo_requisicao"),
        "banco": identificacao.get("banco_comprovante"),
        "fluxo": decisao_fluxo.get("tipo_fluxo"),
        "juros_mora_selic_detectado": _money_value(
            decisao_fluxo.get("juros_mora_selic_detectado")
        ),
        "data_disponibilidade_saque": dados_demo.get("data_disponibilidade_saque"),
        "data_levantamento": dados_comprovante.get("data_levantamento"),
        "mes_recebimento": dados_comprovante.get("mes_recebimento")
        or simulador.get("campo_4_1_mes_recebimento"),
        "numero_meses_rra": _money_value(
            dados_comprovante.get("numero_meses_rra")
            or simulador.get("campo_4_2_numero_meses")
            or receita_resultado.get("numero_meses")
        ),
        "valor_bruto_levantamento": _money_value(
            dados_comprovante.get("valor_bruto_levantamento")
            or simulador.get("campo_3_1_rendimentos_tributaveis")
        ),
        "valor_liquido_levantamento": _money_value(
            dados_comprovante.get("valor_liquido_levantamento")
        ),
        "imposto_retido_na_fonte": _money_value(
            dados_comprovante.get("imposto_retido_na_fonte")
            or imposto_atualizado.get("imposto_retido")
        ),
        "deducao_previdenciaria": _money_value(
            dados_comprovante.get("deducao_previdenciaria")
            or simulador.get("previdencia_oficial")
        ),
        "rendimentos_tributaveis_rra": _money_value(
            receita_resultado.get("rendimentos_tributaveis")
        ),
        "juros_rra": _money_value(
            simulador.get("campo_3_2_valor_dos_juros")
            or receita_resultado.get("juros")
        ),
        "base_calculo_rra": _money_value(receita_resultado.get("base_calculo")),
        "irrf_calculado_receita": _money_value(
            receita_resultado.get("irrf") or imposto_atualizado.get("irrf_calculado")
        ),
        "imposto_indevido": _money_value(imposto_atualizado.get("imposto_indevido")),
        "selic_percentual": _money_value(selic.get("percentual")),
        "selic_acumulado": _money_value(imposto_atualizado.get("selic_acumulado")),
        "valor_final_atualizado": _money_value(
            imposto_atualizado.get("imposto_retido_indevidamente_atualizado")
        ),
        "vencimento_selic": selic.get("vencimento"),
        "pagamento_selic": selic.get("pagamento"),
        "alertas": _join_values(alertas),
        "json_path": str(json_path),
        "demonstrativo_pdf": _first(payload, "arquivos_originais"),
        "comprovante_pdf": _first(payload, "arquivos_originais", index=1),
    }


def _write_summary_sheet(sheet: Any, rows: list[dict[str, Any]]) -> None:
    calculados = [row for row in rows if row["status_final"] == "calculado"]
    nao_calculados = [row for row in rows if row["status_final"] != "calculado"]
    summary_rows = [
        ("Total de registros", len(rows)),
        ("Calculados", len(calculados)),
        ("Não calculados", len(nao_calculados)),
        ("Total imposto retido", _sum_rows(calculados, "imposto_retido_na_fonte")),
        ("Total IRRF calculado Receita", _sum_rows(calculados, "irrf_calculado_receita")),
        ("Total imposto indevido", _sum_rows(calculados, "imposto_indevido")),
        ("Total Selic acumulada", _sum_rows(calculados, "selic_acumulado")),
        ("Total final atualizado", _sum_rows(calculados, "valor_final_atualizado")),
    ]
    sheet.append(["Indicador", "Valor"])
    for label, value in summary_rows:
        sheet.append([label, value])

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row_index in range(5, 10):
        sheet.cell(row=row_index, column=2).number_format = _CURRENCY_FORMAT
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 22


def _write_results_sheet(sheet: Any, rows: list[dict[str, Any]]) -> None:
    headers = [
        "status_final",
        "requerente_nome",
        "cpf_cnpj",
        "processo_originario",
        "processo_requisicao",
        "banco",
        "fluxo",
        "data_disponibilidade_saque",
        "data_levantamento",
        "mes_recebimento",
        "numero_meses_rra",
        "valor_bruto_levantamento",
        "valor_liquido_levantamento",
        "imposto_retido_na_fonte",
        "deducao_previdenciaria",
        "rendimentos_tributaveis_rra",
        "juros_rra",
        "base_calculo_rra",
        "irrf_calculado_receita",
        "imposto_indevido",
        "selic_percentual",
        "selic_acumulado",
        "valor_final_atualizado",
        "rra_status",
        "imposto_atualizado_status",
        "motivo_nao_calculado",
        "vencimento_selic",
        "pagamento_selic",
        "alertas",
        "json_path",
        "demonstrativo_pdf",
        "comprovante_pdf",
    ]
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    currency_columns = {
        "valor_bruto_levantamento",
        "valor_liquido_levantamento",
        "imposto_retido_na_fonte",
        "deducao_previdenciaria",
        "rendimentos_tributaveis_rra",
        "juros_rra",
        "base_calculo_rra",
        "irrf_calculado_receita",
        "imposto_indevido",
        "selic_acumulado",
        "valor_final_atualizado",
    }
    for column_index, header in enumerate(headers, start=1):
        column_letter = get_column_letter(column_index)
        if header in currency_columns:
            for row_index in range(2, sheet.max_row + 1):
                sheet.cell(row=row_index, column=column_index).number_format = (
                    _CURRENCY_FORMAT
                )
        elif header == "selic_percentual":
            for row_index in range(2, sheet.max_row + 1):
                sheet.cell(row=row_index, column=column_index).number_format = "0.00"

        max_length = max(
            len(str(sheet.cell(row=row_index, column=column_index).value or ""))
            for row_index in range(1, sheet.max_row + 1)
        )
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 55)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _get(payload: dict[str, Any], *keys: str) -> Any:
    value: Any = payload
    for key in keys:
        if not isinstance(value, dict):
            return None
        value = value.get(key)
    return value


def _first(payload: dict[str, Any], *keys: str, index: int = 0) -> Any:
    value = _get(payload, *keys)
    if isinstance(value, list) and len(value) > index:
        return value[index]
    return None


def _money_value(value: Any) -> float | None:
    if value is None or value == "":
        return None
    text = str(value).strip().replace("R$", "").strip()
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(Decimal(text))
    except (InvalidOperation, ValueError):
        return None


def _sum_rows(rows: list[dict[str, Any]], key: str) -> float:
    return round(sum(float(row.get(key) or 0) for row in rows), 2)


def _join_values(*values: Any) -> str | None:
    parts: list[str] = []
    for value in values:
        if value is None or value == "":
            continue
        if isinstance(value, list):
            parts.extend(str(item) for item in value if item)
        else:
            parts.append(str(value))
    return "; ".join(parts) if parts else None


_CURRENCY_FORMAT = 'R$ #,##0.00'
